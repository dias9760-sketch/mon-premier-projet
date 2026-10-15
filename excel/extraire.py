"""Lit les onglets mensuels de l'ancien fichier de commandes et renvoie une liste de lignes."""
import datetime as dt
import openpyxl

MOIS = ["JAN", "FEV", "MARS", "AVRIL", "MAI", "JUIN", "JUIL", "AOUT", "SEPT", "OCT", "NOV", "DEC"]
CHAMPS = ["contact", "client", "facture", "date_paiement", "paiement", "numero", "emballage",
          "produit", "couleur", "perso", "prix_vente", "prix_achat", "date_prod", "statut",
          "date_envoi", "recuperation", "adresse"]
HERITES = ["contact", "client", "facture", "date_paiement", "paiement", "numero", "recuperation", "adresse"]


def _propre(v):
    if isinstance(v, str):
        v = v.strip()
        return v or None
    return v


def onglets_mensuels(wb, annee):
    for ws in wb.worksheets:
        nom = ws.title.strip().upper()
        for i, m in enumerate(MOIS):
            if nom == f"{m} {annee}":
                yield i + 1, ws


def extraire(chemin, annee=2026):
    wb = openpyxl.load_workbook(chemin, data_only=True)
    lignes = []
    for mois, ws in onglets_mensuels(wb, annee):
        precedente = None
        for r in range(5, ws.max_row + 1):
            if str(ws.cell(r, 10).value or "").strip().upper() == "TOTAL VENTE":
                break
            v = {c: _propre(ws.cell(r, i + 1).value) for i, c in enumerate(CHAMPS)}
            if not any(v[c] for c in ("produit", "perso", "prix_vente", "client")):
                continue
            # Ligne sans client ni n° = article supplémentaire de la commande précédente
            suite = not v["client"] and not v["facture"] and precedente is not None
            if suite:
                for c in HERITES:
                    if v[c] is None:
                        v[c] = precedente[c]
            v["suite"] = suite
            v["mois"] = mois
            v["onglet"] = ws.title
            v["ligne"] = r
            for c in ("prix_vente", "prix_achat"):
                if not isinstance(v[c], (int, float)):
                    v[c] = None
            for c in ("date_paiement", "date_prod", "date_envoi"):
                if not isinstance(v[c], dt.datetime):
                    v[c] = None
            for c in ("facture", "numero"):
                if isinstance(v[c], float) and v[c].is_integer():
                    v[c] = int(v[c])
            lignes.append(v)
            precedente = v
    return lignes
